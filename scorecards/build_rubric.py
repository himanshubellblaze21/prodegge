"""
Build scorecards/v<N>/rubric.json from the client's Excel templates plus the
hand-written scoring rules in scorecards/v<N>/rules.json.

The Excel template is the source of truth for WHAT is scored — criterion ids,
wording, points, MUST flags and sections — because the client recomputes the
score from it when they open the sheet. rules.json only adds HOW to score
(guidance from the checklist documents, required-phrase checks, durations).
Generating the criteria from the sheet means they can never drift from it.

    python scorecards/build_rubric.py 1            # (re)build v1
    python scorecards/build_rubric.py --check      # every version up to date?

Adding a scorecard version: put the new templates in the repo, add the version
to registry.json (with its templates and S3 keys), write v<N>/rules.json, run
this script, set "current" in registry.json, then deploy. Evaluations record
the version they were scored on, so older ones keep their own template.
"""
import json
import os
import re
import sys

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(HERE)
CALL_TYPES = ('BCM_PHYSICAL_PD', 'BM_AUDIO_FI', 'RCM_AUDIO_PD')


def extract_template(path: str) -> dict:
    ws = openpyxl.load_workbook(path)['Scorecard']
    sections, criteria, current = {}, [], None
    for r in range(1, ws.max_row + 1):
        a = ws.cell(r, 1).value
        if not isinstance(a, str):
            continue
        a = a.strip()
        m = re.match(r'^([A-J])\.\s', a)
        if m:
            current = m.group(1)
            sections[current] = {'pts': ws.cell(r, 3).value, 'title': a,
                                 'time_hint': ws.cell(r, 7).value or ''}
        elif re.match(r'^[A-J]\d{1,2}$', a):
            if current is None or a[0] != current:
                raise ValueError(f'{path}: criterion {a} (row {r}) is outside its section')
            criteria.append({'id': a, 'sec': current, 'pts': ws.cell(r, 3).value,
                             'must': str(ws.cell(r, 4).value or '').strip().upper() == 'MUST',
                             'text': str(ws.cell(r, 2).value or '').strip(), 'row': r})
    return {'sections': sections, 'criteria': criteria}


def build(version: str) -> dict:
    registry = json.load(open(os.path.join(HERE, 'registry.json'), encoding='utf-8'))
    meta = registry['versions'][version]
    rules = json.load(open(os.path.join(HERE, f'v{version}', 'rules.json'), encoding='utf-8'))
    out = {'version': version, 'label': meta['label'], 'client_version': meta['client_version'],
           'call_types': {}}

    for ct in CALL_TYPES:
        tmpl = meta['templates'][ct]
        ext = extract_template(os.path.join(REPO, tmpl['source']))
        r = rules['call_types'][ct]
        crules = r.get('criteria', {})

        unknown = set(crules) - {c['id'] for c in ext['criteria']}
        if unknown:
            raise ValueError(f'v{version} {ct}: rules for criteria not in the template: {sorted(unknown)}')
        missing_names = set(ext['sections']) - set(r['section_names'])
        if missing_names:
            raise ValueError(f'v{version} {ct}: no section name for {sorted(missing_names)}')

        criteria = []
        for c in ext['criteria']:
            extra = crules.get(c['id'], {})
            item = {k: c[k] for k in ('id', 'sec', 'pts', 'must', 'text')}
            if extra.get('guidance'):
                item['guidance'] = extra['guidance']
            if extra.get('absence_counts'):
                item['absence_counts'] = True
            if extra.get('required'):
                item['required'] = {'cap': 'Half', **extra['required']}
            criteria.append(item)

        section_pts = {s: v['pts'] for s, v in ext['sections'].items()}
        for s, pts in section_pts.items():
            got = sum(c['pts'] for c in criteria if c['sec'] == s)
            if got != pts:
                raise ValueError(f'v{version} {ct}: section {s} is {pts} pts but its items add to {got}')
        total = sum(section_pts.values())
        if total != 100:
            raise ValueError(f'v{version} {ct}: sections add to {total}, not 100')

        out['call_types'][ct] = {
            'template_s3_key': tmpl['s3_key'],
            'template_source': tmpl['source'],
            'type_note': r['type_note'],
            'section_pts': section_pts,
            'section_names': {s: r['section_names'][s] for s in section_pts},
            'section_time_hints': {s: v['time_hint'] for s, v in ext['sections'].items()},
            'duration_minutes': r['duration_minutes'],
            'applicable_rule': r['applicable_rule'],
            'e6_input': r['e6_input'],
            'absence_rule': r.get('absence_rule'),
            'general_rules': r.get('general_rules', []),
            'red_flags': r['red_flags'],
            'criteria': criteria,
        }
    return out


def rubric_path(version: str) -> str:
    return os.path.join(HERE, f'v{version}', 'rubric.json')


def render(rubric: dict) -> str:
    return json.dumps(rubric, indent=2, ensure_ascii=False) + '\n'


def main(argv):
    registry = json.load(open(os.path.join(HERE, 'registry.json'), encoding='utf-8'))
    if argv and argv[0] == '--check':
        stale = []
        for v in registry['versions']:
            want = render(build(v))
            have = open(rubric_path(v), encoding='utf-8').read() if os.path.exists(rubric_path(v)) else ''
            if want != have:
                stale.append(v)
        if stale:
            print(f'Out of date: v{", v".join(stale)} — run: python scorecards/build_rubric.py <version>')
            return 1
        print(f'All {len(registry["versions"])} scorecard versions are up to date.')
        return 0

    versions = argv or list(registry['versions'])
    for v in versions:
        rubric = build(v)
        with open(rubric_path(v), 'w', encoding='utf-8', newline='\n') as f:
            f.write(render(rubric))
        for ct, cfg in rubric['call_types'].items():
            musts = [c['id'] for c in cfg['criteria'] if c['must']]
            print(f'v{v} {ct}: {len(cfg["criteria"])} items, sections {cfg["section_pts"]}, '
                  f'{len(musts)} MUST ({", ".join(musts)})')
    return 0


if __name__ == '__main__':
    sys.exit(main(sys.argv[1:]))
