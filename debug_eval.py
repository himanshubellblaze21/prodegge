import boto3, json
from decimal import Decimal

dynamodb = boto3.resource('dynamodb', region_name='ap-south-1')
s3 = boto3.client('s3', region_name='ap-south-1')
table = dynamodb.Table('audio-pd-evaluations-dev')

# Scan for eval-6ef36091fa63 — find ALL records with this ID
result = table.scan(
    FilterExpression='evaluation_id = :eid',
    ExpressionAttributeValues={':eid': 'eval-6ef36091fa63'}
)
items = result['Items']
print(f"DynamoDB records for eval-6ef36091fa63: {len(items)}")
for item in items:
    print(f"\n  evaluation_id : {item.get('evaluation_id')}")
    print(f"  created_at    : {item.get('created_at')}")
    print(f"  application_id: {item.get('application_id')}")
    print(f"  total_score   : {item.get('total_score')}")
    print(f"  grade_band    : {item.get('grade_band')}")
    print(f"  status        : {item.get('status')}")
    print(f"  result_s3_key : {item.get('result_s3_key')}")
    print(f"  excel_s3_key  : {item.get('excel_s3_key')}")

# Also check what score is in the evaluation-result.json
print("\n--- Checking S3 evaluation result ---")
try:
    resp = s3.get_object(
        Bucket='audio-pd-reports-dev',
        Key='evaluations/eval-6ef36091fa63/evaluation-result.json'
    )
    data = json.loads(resp['Body'].read())
    scoring = data.get('scoring', {})
    print(f"  S3 result total_score: {scoring.get('total_score')}")
    print(f"  S3 result grade_band : {scoring.get('grade_band')}")
    print(f"  S3 result verdict    : {scoring.get('verdict')}")
    print(f"  application_id       : {data.get('application_id')}")
    print(f"  evaluated_at         : {data.get('evaluated_at')}")
    
    # Check criteria points
    criteria = data.get('ai_evaluation', {}).get('criteria', [])
    print(f"\n  Criteria count: {len(criteria)}")
    total_pts = sum(float(c.get('points_scored', 0)) for c in criteria)
    print(f"  Sum of points_scored: {total_pts}")
    
    # Check what the derived compute gives
    try:
        import sys
        import os
        sys.path.append(os.path.join(os.getcwd(), 'lambda', 'evaluation'))
        from handler import compute_derived
        derived = compute_derived(data)
        print(f"\n  Computed score (Python): {derived['total_score']}")
    except Exception as ex:
        print(f"  Could not compute derived score: {ex}")
except Exception as e:
    print(f"  Error: {e}")

# Check what the Excel shows
print("\n--- Checking Excel scorecard ---")
try:
    import openpyxl
    from io import BytesIO
    resp2 = s3.get_object(
        Bucket='audio-pd-reports-dev',
        Key='evaluations/eval-6ef36091fa63/scorecard.xlsx'
    )
    wb = openpyxl.load_workbook(BytesIO(resp2['Body'].read()))
    ws = wb['Scorecard']
    print(f"  Excel E72 (Total Score): {ws['E72'].value}")
    print(f"  Excel E73 (Knockout):    {ws['E73'].value}")
    print(f"  Excel E74 (Grade Band):  {ws['E74'].value}")
    print(f"  Excel E75 (Verdict):     {ws['E75'].value}")
    print(f"  Excel B2 (App ID):       {ws['B2'].value}")
    # Show H column (points per criterion)
    print("\n  H column (points per criterion):")
    for row in range(11, 49):
        a = ws[f'A{row}'].value
        h = ws[f'H{row}'].value
        g = ws[f'G{row}'].value
        if a and str(a).strip() and not str(a).startswith('SECTION'):
            print(f"    {a}: G={g} H={h}")
except Exception as e:
    print(f"  Error: {e}")
