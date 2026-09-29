"""
Fast offline tests for the evaluation Lambda's safety nets. No AWS calls, no
Bedrock, no cost — run this on every change:

    python tests/test_safety_nets.py

These cover the failure modes that previously produced silently-wrong
scorecards: unparseable model output, non-canonical score words, hallucinated
criterion ids, and transient Bedrock failures.
"""
import io
import os
import sys

# Devanagari in test labels would crash the default Windows console codec.
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

os.environ.setdefault('TRANSCRIPTS_BUCKET', 'test-transcripts')
os.environ.setdefault('REPORTS_BUCKET', 'test-reports')
os.environ.setdefault('DYNAMODB_TABLE', 'test-table')
os.environ.setdefault('AWS_REGION', 'ap-south-1')
sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', 'lambda', 'evaluation'))

import handler  # noqa: E402

passed = failed = 0


def check(label, got, want):
    global passed, failed
    if got == want:
        passed += 1
        print(f"  ok   {label}")
    else:
        failed += 1
        print(f"  FAIL {label}: got {got!r}, want {want!r}")


print("normalise_score — model synonyms must not silently score zero")
for raw, want in [
    ('Yes', 'Yes'), ('yes', 'Yes'), ('YES', 'Yes'), ('Full', 'Yes'), ('fully covered', 'Yes'),
    ('Half', 'Half'), ('partial', 'Half'), ('Partially Covered', 'Half'),
    ('No', 'No'), ('none', 'No'), ('not covered', 'No'),
    ('N/A', 'N/A'), ('na', 'N/A'), ('Not Applicable', 'N/A'),
    ('  Yes  ', 'Yes'),
    ('gibberish', 'No'), (None, 'No'), (42, 'No'),
]:
    check(f"normalise_score({raw!r})", handler.normalise_score(raw), want)

print("\n_extract_json_object — tolerate prose/fences around the JSON")
check("plain json", handler._extract_json_object('{"a": 1}'), {'a': 1})
check("prose wrapped",
      handler._extract_json_object('Here is the result:\n{"a": 1}\nHope that helps!'), {'a': 1})
check("nested braces",
      handler._extract_json_object('noise {"a": {"b": 2}} trailing'), {'a': {'b': 2}})
try:
    handler._extract_json_object('no json at all')
    check("raises on garbage", 'no raise', 'ValueError')
except ValueError:
    check("raises on garbage", 'ValueError', 'ValueError')

print("\ncall_bedrock_json — retries transient failures instead of defaulting to zero")
calls = {'n': 0}
original = handler.call_bedrock


def flaky(prompt, label):
    calls['n'] += 1
    if calls['n'] < 3:
        raise Exception("throttled")
    return '{"items": [{"id": "A1", "score": "Yes"}]}'


handler.call_bedrock = flaky
check("succeeds on 3rd attempt",
      handler.call_bedrock_json('p', 'TEST')['items'][0]['score'], 'Yes')
check("made 3 attempts", calls['n'], 3)

calls['n'] = 0


def always_bad(prompt, label):
    calls['n'] += 1
    return 'not json'


handler.call_bedrock = always_bad
try:
    handler.call_bedrock_json('p', 'TEST', attempts=2)
    check("raises after exhausting attempts", 'no raise', 'Exception')
except Exception:
    check("raises after exhausting attempts", 'Exception', 'Exception')
check("respected attempts=2", calls['n'], 2)
handler.call_bedrock = original

print("\ncalculate_scores — must match the client's Excel template formulas exactly")
# Template: applicable = 100 (BM/RCM), scored = SUM(Yes->pts, Half->pts/2),
# SCORE = ROUND(scored/applicable*100, 0). There is no per-item N/A in the
# sheet — an N/A mark scores zero and still counts in the denominator.
criteria, _, _ = handler.get_config('BM_AUDIO_FI')
first_three = criteria[:3]
items = [
    {'id': first_three[0]['id'], 'score': 'Yes'},
    {'id': first_three[1]['id'], 'score': 'Half'},
    {'id': first_three[2]['id'], 'score': 'N/A'},
]
items += [{'id': c['id'], 'score': 'No'} for c in criteria[3:]]
scoring = handler.calculate_scores({'items': items, '_duration_secs': 600}, 'BM_AUDIO_FI')
check("BM applicable is always 100", scoring['total_applicable'], 100.0)
expected_scored = first_three[0]['pts'] + first_three[1]['pts'] / 2.0
check("Yes + Half scored correctly", scoring['total_scored'], float(expected_scored))
check("score is rounded to whole number",
      scoring['total_score'], float(round(expected_scored / 100 * 100)))

print("\ncalculate_scores — BCM section C: legacy drops it without other income, v1 never does")
bcm_criteria, _, _ = handler.get_config('BCM_PHYSICAL_PD', '0')
bcm_all_no = [{'id': c['id'], 'score': 'No'} for c in bcm_criteria]
s_with = handler.calculate_scores({'items': bcm_all_no, '_duration_secs': 1200},
                                  'BCM_PHYSICAL_PD', other_income=True, version='0')
s_without = handler.calculate_scores({'items': bcm_all_no, '_duration_secs': 1200},
                                     'BCM_PHYSICAL_PD', other_income=False, version='0')
check("legacy: with other income -> 100 applicable", s_with['total_applicable'], 100.0)
check("legacy: without other income -> 90 applicable", s_without['total_applicable'], 90.0)
v1_bcm, _, _ = handler.get_config('BCM_PHYSICAL_PD', '1')
s_v1 = handler.calculate_scores({'items': [{'id': c['id'], 'score': 'No'} for c in v1_bcm],
                                 '_duration_secs': 1200}, 'BCM_PHYSICAL_PD',
                                other_income=False, version='1')
check("v1: Group C always scored -> 100 applicable", s_v1['total_applicable'], 100.0)
check("v1: result records its scorecard version", s_v1['scorecard_version'], '1')
check("BM is never reduced",
      handler.calculate_scores({'items': items, '_duration_secs': 600}, 'BM_AUDIO_FI',
                               other_income=False)['total_applicable'], 100.0)

print("\nhas_other_income — drives the template's E6 cell")
check("yes on dairy", handler.has_other_income({'has_dairy_or_livestock': 'yes'}), True)
check("yes on farm land", handler.has_other_income({'has_farm_land': 'yes'}), True)
check("no when all absent",
      handler.has_other_income({'has_farm_land': 'no', 'has_dairy_or_livestock': 'no'}), False)
check("unknown profile keeps section C", handler.has_other_income({}), True)
check("unclear keeps section C",
      handler.has_other_income({'has_farm_land': 'unclear',
                                'has_dairy_or_livestock': 'unclear'}), True)

print("\ncalculate_scores — MUST is 'not exactly Yes' (template counts Half and N/A as missed)")
all_yes = [{'id': c['id'], 'score': 'Yes'} for c in criteria]
scoring = handler.calculate_scores({'items': all_yes, '_duration_secs': 600}, 'BM_AUDIO_FI')
check("all Yes -> EXCELLENT", scoring['grade'], 'EXCELLENT')
first_must = next(c for c in criteria if c['must'])
for bad in ('Half', 'No', 'N/A'):
    downgraded = [{'id': c['id'], 'score': (bad if c['id'] == first_must['id'] else 'Yes')}
                  for c in criteria]
    s = handler.calculate_scores({'items': downgraded, '_duration_secs': 600}, 'BM_AUDIO_FI')
    check(f"one MUST at {bad} -> NOT ACCEPTED", s['grade'], 'NOT ACCEPTED')

print("\ncalculate_scores — duration grace window")
# BM nominal window is 5-15 min; the grace window must accept a 15:36 call that
# the old hard cutoff rejected outright.
scoring = handler.calculate_scores({'items': all_yes, '_duration_secs': 936}, 'BM_AUDIO_FI')
check("15:36 BM call accepted", scoring['recording_invalid'], False)
scoring = handler.calculate_scores({'items': all_yes, '_duration_secs': 60}, 'BM_AUDIO_FI')
check("1 min BM call still rejected", scoring['recording_invalid'], True)

print("\nordered_sections — every criterion belongs to a returned section")
for ct in ('BCM_PHYSICAL_PD', 'BM_AUDIO_FI', 'RCM_AUDIO_PD'):
    crit, section_pts, section_names = handler.get_config(ct)
    secs = handler.ordered_sections(crit)
    check(f"{ct}: sections match section_pts keys", sorted(secs), sorted(section_pts.keys()))
    check(f"{ct}: every section has a name", all(s in section_names for s in secs), True)
    check(f"{ct}: section points sum to 100", sum(section_pts.values()), 100)
    per_section = {}
    for c in crit:
        per_section[c['sec']] = per_section.get(c['sec'], 0) + c['pts']
    check(f"{ct}: criteria points match section_pts", per_section, section_pts)

print("\nformat_timestamp")
check("0s", handler.format_timestamp(0), '00:00')
check("125.7s", handler.format_timestamp(125.7), '02:05')
check("3661s", handler.format_timestamp(3661), '61:01')
check("garbage", handler.format_timestamp('abc'), '')
check("None", handler.format_timestamp(None), '')

print("\nverify_evidence — timestamps come from segment data, never the model")
SEGS = [
    {'start': 0.7,   'end': 4.7,   'speaker': 'spk_0', 'text': 'नमस्कार मेरा नाम राहुल राजपूत है'},
    {'start': 21.8,  'end': 25.0,  'speaker': 'spk_1', 'text': 'गाय भैंस'},
    {'start': 169.3, 'end': 174.0, 'speaker': 'spk_0', 'text': 'ये पिताजी के नाम से है जी हाँ'},
    {'start': 200.0, 'end': 205.0, 'speaker': 'spk_1', 'text': 'गाय भैंस'},
    {'start': 320.5, 'end': 326.0, 'speaker': 'spk_0', 'text': 'दो भैंस तीन गाय हैं चार हो जाएंगे'},
]

r = handler.verify_evidence('ये पिताजी के नाम से है', SEGS)
check("exact quote -> real segment start", r['timestamp'], '02:49')
check("exact quote -> verified", r['verified'], True)
check("exact quote -> speaker carried", r['speaker'], 'spk_0')

r = handler.verify_evidence('दो भैंस तीन गाय हैं चार हो जाएंगे', SEGS)
check("second exact quote -> right timestamp", r['timestamp'], '05:20')

# A fabricated quote must never receive a timestamp — that is the whole point.
r = handler.verify_evidence('मैंने चाँद पर दुकान देखी और हाथी गिने वहाँ पर', SEGS)
check("fabricated quote -> not verified", r['verified'], False)
check("fabricated quote -> no timestamp", r['timestamp'], '')

# Too-short quotes are ambiguous and must not resolve to an arbitrary moment.
r = handler.verify_evidence('गाय भैंस', SEGS)
check("2-word quote -> rejected as too short", r['match'], 'too-short')
check("2-word quote -> not verified", r['verified'], False)
check("empty quote -> not verified", handler.verify_evidence('', SEGS)['verified'], False)
check("no segments -> not verified", handler.verify_evidence('कुछ भी लंबा वाक्य यहाँ', [])['verified'], False)

# Repeated phrases are flagged so a reviewer knows the location isn't unique.
r = handler.verify_evidence('नमस्कार मेरा नाम राहुल राजपूत है', SEGS)
check("unique quote reports 'exact'", r['match'], 'exact')

# Punctuation/casing differences must not break matching.
r = handler.verify_evidence('ये पिताजी के नाम से है।', SEGS)
check("punctuation tolerated", r['timestamp'], '02:49')

# A quote spanning consecutive segments resolves to where it starts.
r = handler.verify_evidence('ये पिताजी के नाम से है जी हाँ गाय भैंस', SEGS)
check("spanning quote -> start segment", r['timestamp'], '02:49')

# The evaluator sometimes pastes a long passage spanning many segments instead
# of the short quote it was asked for (seen in production: a 162-token dump).
# It must still resolve — to where the passage BEGINS — rather than being
# written off as unverifiable. Built here from real consecutive segment text,
# the way such a dump actually looks, plus trailing drift.
long_quote = ' '.join([SEGS[2]['text'], SEGS[3]['text'], SEGS[4]['text']]) \
    + ' इसके बाद बातचीत जारी रही और अन्य विषयों पर चर्चा हुई थी'
r = handler.verify_evidence(long_quote, SEGS)
check("long pasted passage -> anchors on its opening", r['timestamp'], '02:49')
check("long pasted passage -> verified", r['verified'], True)

# A long passage of pure invention must still be rejected.
r = handler.verify_evidence('चाँद पर हाथी गिने वहाँ ' * 15, SEGS)
check("long fabricated passage -> rejected", r['verified'], False)

print("\nbuild_timestamped_transcript — the summary's [MM:SS] must be real")
tt = handler.build_timestamped_transcript(SEGS)
check("first line carries real timestamp", tt.splitlines()[0].startswith('[00:00]'), True)
check("includes the 02:49 line", '[02:49]' in tt, True)
check("includes speaker labels", 'spk_0:' in tt, True)
check("falls back when no segments", handler.build_timestamped_transcript([], 'plain'), 'plain')

print("\nrequired utterances — criteria that hinge on a specific thing being said")
# A transcript containing every required utterance must cap nothing. This is the
# guard against a broken or over-tight phrase list silently failing every
# recording, which is far worse than missing a cap.
RICH = ('यह कॉल रिकॉर्ड हो रही है। लॉगिन फीस ली गयी। कमीशन किसी ने नहीं माँगा, कोई ब्रोकर नहीं। '
        'प्रॉपर्टी गिरवी रखी जाएगी। केसीसी चल रहा है। अकाउंट एग्रीगेटर चेक किया। पड़ोसी से बात की। '
        'स्टॉक गिना। फोटो अपलोड की। सेल डीड और पट्टा देखा। कोर्ट केस नहीं है। नाच मैंडेट पर सहमति। '
        'क्लीन ट्रैक रिवॉर्ड बताया। सैंक्शन केएफएस डिस्बर्समेंट बताया। फाइल फॉरवर्ड करता हूँ, अप्रूव। '
        'यह फ्रेश लोन है, कोई बीटी नहीं।')
for ver in ('0', '1'):
    for ct in ('BCM_PHYSICAL_PD', 'BM_AUDIO_FI', 'RCM_AUDIO_PD'):
        crit, _, _ = handler.get_config(ct, ver)
        rules = handler.required_utterances(ct, ver)
        check(f"v{ver} {ct}: has rules", len(rules) > 0, True)
        check(f"v{ver} {ct}: every rule id is a real criterion",
              sorted(rules) == sorted(set(rules) & {c['id'] for c in crit}), True)
        items = [{'id': c['id'], 'score': 'Yes', 'note': ''} for c in crit]
        check(f"v{ver} {ct}: rich transcript caps nothing",
              handler.enforce_required_utterances(RICH, items, ct, ver), [])
        # An empty transcript must cap every rule — proving each one can fire.
        items = [{'id': c['id'], 'score': 'Yes', 'note': ''} for c in crit]
        fired = handler.enforce_required_utterances('कुछ और बात', items, ct, ver)
        check(f"v{ver} {ct}: every rule can fire", sorted(fired), sorted(rules))

# Caps must never raise a mark, and never lower one that is already lower.
items = [{'id': 'A1', 'score': 'No', 'note': 'x'}, {'id': 'H3', 'score': 'Half', 'note': 'x'}]
handler.enforce_required_utterances('कुछ और बात', items, 'RCM_AUDIO_PD')
check("a 'No' is not raised to the cap", items[0]['score'], 'No')
check("an existing 'Half' is left alone", items[1]['score'], 'Half')

# Transcribe mangles Hindi: "गिरवी" came through as "गिरी हुई" on a real call,
# and matching the full word would have invented a failure.
items = [{'id': 'B4', 'score': 'Yes', 'note': ''}]
handler.enforce_required_utterances('आपकी प्रॉपर्टी गिरी हुई रखी जाएगी', items, 'RCM_AUDIO_PD')
check("mis-transcribed 'गिरवी' still counts as said", items[0]['score'], 'Yes')

print("\nMUST flags must match the client templates' own MUST-count formulas")
# Read straight from the =IF(E12<>"Yes",1,0)+... formula in each version's
# Scorecard. That formula drives the sheet's own RESULT, so a mismatch means our
# verdict and the client's sheet disagree about whether a recording is
# acceptable. (We once had 16 MUST items on BCM against the template's 9.)
import re as _re
import openpyxl as _opx
sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..'))
import scorecards as _sc
for ver in _sc.registry()['versions']:
    for ct in ('BCM_PHYSICAL_PD', 'BM_AUDIO_FI', 'RCM_AUDIO_PD'):
        src = _sc.config(ct, ver)['template_source']
        ws = _opx.load_workbook(os.path.join(os.path.dirname(__file__), '..', src))['Scorecard']
        formula = next(ws.cell(r, 3).value for r in range(1, ws.max_row + 1)
                       if str(ws.cell(r, 1).value or '').lower().startswith('must items not fully covered'))
        rows = [int(x) for x in _re.findall(r'IF\(E(\d+)<>"Yes"', formula)]
        template_must = sorted(str(ws.cell(r, 1).value).strip() for r in rows)
        crit, _, _ = handler.get_config(ct, ver)
        check(f"v{ver} {ct}: MUST set matches the template formula",
              sorted(c['id'] for c in crit if c['must']), template_must)
check("v1 MUST counts match the V2.1 rulebooks (BCM 9 / BM 12 / RCM 11)",
      [sum(c['must'] for c in handler.get_config(ct, '1')[0])
       for ct in ('BCM_PHYSICAL_PD', 'BM_AUDIO_FI', 'RCM_AUDIO_PD')], [9, 12, 11])

print("\nExcel value caching — the score must be readable without recalculating")
# A formula written by openpyxl has no cached result, so every viewer that does
# not recalculate shows a BLANK score. Values are injected alongside the
# formulas so the sheet reads correctly everywhere and still stays live.
try:
    import importlib.util as _ilu
    import zipfile as _zf
    _xl_path = os.path.join(os.path.dirname(__file__), '..', 'lambda', 'excel-generator')
    sys.path.insert(0, os.path.join(_xl_path, 'package'))
    os.environ.setdefault('REPORTS_BUCKET', 'test'), os.environ.setdefault('DYNAMODB_TABLE', 'test')
    _spec = _ilu.spec_from_file_location('excel_handler2', os.path.join(_xl_path, 'handler.py'))
    _eh = _ilu.module_from_spec(_spec)
    _spec.loader.exec_module(_eh)
    import openpyxl as _op
    from io import BytesIO as _BIO

    _wb = _op.Workbook()
    _ws = _wb.active
    _ws.title = 'Scorecard'
    _ws['C1'] = 5
    _ws['C2'] = '=C1*2'
    _ws['C3'] = '=IF(C1>0,"GOOD","BAD")'
    _buf = _BIO()
    _wb.save(_buf)
    _out = _eh.cache_formula_values(_buf.getvalue(), 'Scorecard', {'C2': 10, 'C3': 'GOOD'})
    _vals = _op.load_workbook(_BIO(_out), data_only=True)['Scorecard']
    _forms = _op.load_workbook(_BIO(_out))['Scorecard']
    check("numeric formula has a cached value", _vals['C2'].value, 10)
    check("text formula has a cached value", _vals['C3'].value, 'GOOD')
    check("numeric formula preserved", _forms['C2'].value, '=C1*2')
    check("text formula preserved", _forms['C3'].value, '=IF(C1>0,"GOOD","BAD")')
    check("untouched cell unaffected", _vals['C1'].value, 5)
    try:
        _eh.cache_formula_values(_buf.getvalue(), 'NoSuchSheet', {'C2': 1})
        check("missing sheet raises", 'no raise', 'KeyError')
    except KeyError:
        check("missing sheet raises", 'KeyError', 'KeyError')
except Exception as _e:  # pragma: no cover
    print(f"  skip (excel handler not importable here: {_e})")

print("\nCross-Lambda consistency — both handlers read the same scorecard registry")
# lambda/evaluation and lambda/excel-generator each load the rubric of a given
# version. If they ever disagree, the Excel silently recomputes different
# scores from the same items than the evaluator reported.
sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', 'lambda',
                                'excel-generator', 'package'))
sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', 'lambda', 'excel-generator'))
try:
    import importlib.util
    spec = importlib.util.spec_from_file_location(
        'excel_handler',
        os.path.join(os.path.dirname(__file__), '..', 'lambda', 'excel-generator', 'handler.py'))
    excel_handler = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(excel_handler)
except Exception as e:  # pragma: no cover - environment without openpyxl
    print(f"  skip (excel handler not importable here: {e})")
    excel_handler = None

if excel_handler is not None:
    for ver in ('0', '1'):
        for ct in ('BCM_PHYSICAL_PD', 'BM_AUDIO_FI', 'RCM_AUDIO_PD'):
            eval_crit, eval_pts, eval_names = handler.get_config(ct, ver)
            xl_crit, xl_pts, xl_names = excel_handler.get_config(ct, ver)
            check(f"v{ver} {ct}: same criterion ids",
                  [c['id'] for c in eval_crit], [c['id'] for c in xl_crit])
            check(f"v{ver} {ct}: same points per criterion",
                  [c['pts'] for c in eval_crit], [c['pts'] for c in xl_crit])
            check(f"v{ver} {ct}: same MUST flags",
                  [c['must'] for c in eval_crit], [c['must'] for c in xl_crit])
            check(f"v{ver} {ct}: same section points", eval_pts, xl_pts)
            check(f"v{ver} {ct}: same section names", eval_names, xl_names)
    check("score normalisation agrees across Lambdas",
          [excel_handler.normalise_score(v) for v in ('Full', 'partial', 'na', 'weird')],
          [handler.normalise_score(v) for v in ('Full', 'partial', 'na', 'weird')])

print("\nScorecard registry")
import subprocess as _sp
_r = _sp.run([sys.executable, os.path.join(os.path.dirname(__file__), '..', 'scorecards', 'build_rubric.py'),
              '--check'], capture_output=True, text=True)
check("every rubric.json is rebuilt from its templates + rules", _r.returncode, 0)
check("current version is v1", _sc.current_version(), '1')
check("unrecorded version resolves to legacy", _sc.resolve_version(None), '0')
check("unknown version resolves to legacy", _sc.resolve_version('99'), '0')
check("recorded version kept", _sc.resolve_version('1'), '1')
for ver in _sc.registry()['versions']:
    for ct in ('BCM_PHYSICAL_PD', 'BM_AUDIO_FI', 'RCM_AUDIO_PD'):
        names = {c['required']['phrases'] for c in handler.get_config(ct, ver)[0] if c.get('required')}
        check(f"v{ver} {ct}: every required phrase set exists", sorted(names - set(handler.PHRASE_SETS)), [])

print("\nv1 prompt rules — stated absence counts, silence does not")
crit, _, names = handler.get_config('BCM_PHYSICAL_PD', '1')
sec_c = [c for c in crit if c['sec'] == 'C']
p1 = handler.build_section_prompt('T', 'BCM_PHYSICAL_PD', 'C', sec_c, names['C'], 1200, {}, '1')
p0 = handler.build_section_prompt('T', 'BCM_PHYSICAL_PD', 'C',
                                  [c for c in handler.get_config('BCM_PHYSICAL_PD', '0')[0] if c['sec'] == 'C'],
                                  names['C'], 1200, {}, '0')
check("v1 prompt names the V2.1 framework", 'V2.1 Simple framework' in p1, True)
check("v1 prompt carries the absence rule", 'SAYS SO clearly' in p1, True)
check("v1 prompt flags absence-eligible criteria", p1.count('counts as fully covered → Yes'), len(sec_c))
check("v1 prompt carries per-criterion guidance", 'How to score:' in p1, True)
check("legacy prompt keeps the V2.0 wording", 'V2.0 Simple framework' in p0 and 'SAYS SO' not in p0, True)

# Re-check: a verified explicit statement of absence goes to Yes only on an
# absence-eligible criterion; anything else found is still capped at Half.
SEGS_ABS = [{'start': 30.0, 'end': 35.0, 'speaker': 'spk_0',
             'text': 'इनके पास कोई खेती की ज़मीन नहीं है और कोई पशु भी नहीं हैं'},
            {'start': 60.0, 'end': 65.0, 'speaker': 'spk_0',
             'text': 'घर का बिजली बिल इनके पिताजी के नाम पर आता है'}]
_orig = handler.call_bedrock_json
handler.call_bedrock_json = lambda prompt, label, attempts=3: {'findings': [
    {'id': 'C1', 'found': True, 'stated_absence': True,
     'evidence': 'इनके पास कोई खेती की ज़मीन नहीं है'},
    {'id': 'E1', 'found': True, 'stated_absence': True,
     'evidence': 'घर का बिजली बिल इनके पिताजी के नाम पर आता है'}]}
try:
    items = [{'id': 'C1', 'score': 'No'}, {'id': 'E1', 'score': 'No'}]
    handler.recheck_no_items('T', SEGS_ABS, items, crit,
                             absence_rule=handler.get_rubric('BCM_PHYSICAL_PD', '1')['absence_rule'])
    check("stated absence on an eligible item -> Yes", items[0]['score'], 'Yes')
    check("absence claim on a non-eligible item -> Half", items[1]['score'], 'Half')
    items = [{'id': 'C1', 'score': 'No'}]
    handler.recheck_no_items('T', SEGS_ABS, items, handler.get_config('BCM_PHYSICAL_PD', '0')[0])
    check("legacy re-check never goes above Half", items[0]['score'], 'Half')
finally:
    handler.call_bedrock_json = _orig

ent = [{'id': 'C1', 'text': 'LAND', 'score': 'Yes', 'evidence': 'x', 'absence_counts': True}]
check("v1 audit accepts absence quotes",
      'does NOT exist' in handler.build_relevance_audit_prompt('T', ent, 'rule'), True)
check("legacy audit prompt unchanged by the absence rule",
      'does NOT exist' in handler.build_relevance_audit_prompt('T', ent), False)

print(f"\n{passed} passed, {failed} failed")
sys.exit(1 if failed else 0)
